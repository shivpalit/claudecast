"""
Orchestration layer — content generation pipeline.
"""

import json
from datetime import datetime
from pathlib import Path

from .config import preferences_path, resolve_config


def _read_input(source: str) -> str:
    """Read input from a file path, '-' (stdin), or treat as literal text."""
    if source == "-":
        import sys
        return sys.stdin.read()

    p = Path(source)
    if p.exists():
        suffix = p.suffix.lower()
        if suffix in (".txt", ".md"):
            return p.read_text()
        if suffix == ".pdf":
            import pdfplumber
            with pdfplumber.open(p) as pdf:
                return "\n\n".join(page.extract_text() or "" for page in pdf.pages)
        if suffix == ".docx":
            from docx import Document
            doc = Document(p)
            return "\n\n".join(para.text for para in doc.paragraphs if para.text.strip())
        if suffix == ".csv":
            import pandas as pd
            df = pd.read_csv(p)
            return df.to_string(index=False)
        return p.read_text()

    return source


def _build_system_prompt(project: str | None) -> str:
    """Concatenate global + project preferences, stripping comment headers."""
    parts = []
    for path in [preferences_path(None), preferences_path(project) if project else None]:
        if path and path.exists():
            text = path.read_text()
            marker = "<!-- preferences below this line -->"
            if marker in text:
                text = text.split(marker, 1)[1]
            text = text.strip()
            if text:
                parts.append(text)
    return "\n\n".join(parts)


def _output_dir(base: str, project: str | None) -> Path:
    stamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    folder = project or "default"
    out = Path(base) / folder / stamp
    out.mkdir(parents=True, exist_ok=True)
    return out


def generate_slides(
    source: str,
    *,
    project: str | None = None,
    output_base: str | None = None,
    slide_count: int | None = None,
) -> dict:
    """
    Slides-only pipeline: input → slide agent → PPTX + slide JSON.
    Returns dict with keys: slides, pptx_path, output_dir.
    """
    from .agents import run_slide_agent

    cfg = resolve_config(project)
    _base = output_base or cfg.get("output_dir", str(Path.home() / "claudecast-output"))

    input_text = _read_input(source)
    system_prompt = _build_system_prompt(project)

    out = _output_dir(_base, project)

    result = run_slide_agent(input_text, str(out), system_prompt, slide_count)
    result["output_dir"] = str(out)

    (out / "manifest.json").write_text(json.dumps(result, indent=2))
    return result


def generate_video(
    source: str,
    *,
    project: str | None = None,
    output_base: str | None = None,
    slide_count: int | None = None,
    voice: str | None = None,
) -> dict:
    """
    Full pipeline: input → slides → voiceover scripts → audio.
    Returns dict with all artifact paths.
    """
    from .agents import run_script_agent, run_slide_agent
    from .compilers.audio import combine_audio, generate_all
    from .compilers.slides import render_slide_images
    from .compilers.video import build_video

    cfg = resolve_config(project)
    _voice = voice or cfg.get("default_voice", "en-US-AriaNeural")
    _base = output_base or cfg.get("output_dir", str(Path.home() / "claudecast-output"))

    input_text = _read_input(source)
    system_prompt = _build_system_prompt(project)

    out = _output_dir(_base, project)

    result = run_slide_agent(input_text, str(out), system_prompt, slide_count)

    print("generating voiceover scripts...")
    scripts = run_script_agent(result["slides"], system_prompt)

    audio_dir = out / "audio"
    print(f"generating audio ({_voice})...")
    audio_paths = generate_all(scripts, str(audio_dir), _voice)

    print("rendering slide images...")
    image_paths = render_slide_images(result["pptx_path"], str(out / "images"))

    combined_path = str(out / "combined.mp3")
    print("combining audio...")
    combine_audio(audio_paths, combined_path)

    video_path = str(out / "video.mp4")
    print("building video...")
    build_video(image_paths, audio_paths, video_path)

    final = {
        "slides": result["slides"],
        "pptx_path": result["pptx_path"],
        "scripts": scripts,
        "audio_paths": audio_paths,
        "combined": combined_path,
        "video_path": video_path,
        "output_dir": str(out),
    }

    (out / "manifest.json").write_text(json.dumps(final, indent=2))
    return final


def generate_podcast(
    source: str,
    *,
    project: str | None = None,
    output_base: str | None = None,
    voice: str | None = None,
) -> dict:
    """
    Podcast pipeline: input → single narration script → one MP3.
    Returns dict with keys: script, audio_path, output_dir.
    """
    from .agents import run_podcast_agent
    from .compilers.audio import generate_slide_audio

    cfg = resolve_config(project)
    _voice = voice or cfg.get("default_voice", "en-US-AriaNeural")
    _base = output_base or cfg.get("output_dir", str(Path.home() / "claudecast-output"))

    input_text = _read_input(source)
    system_prompt = _build_system_prompt(project)

    print("generating narration script...")
    script = run_podcast_agent(input_text, system_prompt)

    out = _output_dir(_base, project)
    audio_path = str(out / "podcast.mp3")

    print(f"generating audio ({_voice})...")
    generate_slide_audio(script, _voice, audio_path)

    result = {
        "script": script,
        "audio_path": audio_path,
        "output_dir": str(out),
    }

    (out / "manifest.json").write_text(json.dumps(result, indent=2))
    return result
